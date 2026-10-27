"""Application tracker with duplicate-application protection and exports."""
from __future__ import annotations

import csv
import io
import json
from datetime import datetime, timedelta, timezone
from typing import Any

from . import db
from .audit_log import log
from .models import TRACKER_STATUSES

APPLIED_LIKE = {"Applied", "Email Sent", "Follow-up Due", "Interview", "Assessment", "Offer Received"}
FOLLOWUP_DAYS = 7


class DuplicateApplicationError(Exception):
    pass


def _job(job_id: str) -> dict[str, Any] | None:
    return db.query_one("SELECT * FROM jobs WHERE id=?", (job_id,))


def get_for_job(job_id: str) -> dict[str, Any] | None:
    row = db.query_one("SELECT * FROM applications WHERE job_id=? ORDER BY attempt DESC LIMIT 1", (job_id,))
    return _decorate(row) if row else None


def _decorate(row: dict[str, Any]) -> dict[str, Any]:
    row["history"] = db.jload(row["history"], [])
    if row["status"] in ("Applied", "Email Sent") and row.get("followup_due"):
        if row["followup_due"] <= datetime.now(timezone.utc).isoformat():
            row["status_display"] = "Follow-up Due"
    row.setdefault("status_display", row["status"])
    return row


def _related_applied(job: dict[str, Any]) -> dict[str, Any] | None:
    """Same role reached via another source (same fingerprint) already applied?"""
    return db.query_one(
        "SELECT a.* FROM applications a JOIN jobs j ON j.id=a.job_id WHERE j.fingerprint=? AND a.job_id!=? AND a.status IN (%s)"
        % ",".join("?" * len(APPLIED_LIKE)),
        (job["fingerprint"], job["id"], *APPLIED_LIKE),
    )


def set_status(
    job_id: str,
    status: str,
    *,
    notes: str | None = None,
    method: str | None = None,
    confirmation: str | None = None,
    reapply: bool = False,
) -> dict[str, Any]:
    if status not in TRACKER_STATUSES:
        raise ValueError(f"Unknown status {status}")
    job = _job(job_id)
    if not job:
        raise KeyError("Job not found")
    now = db.now_iso()
    current = db.query_one("SELECT * FROM applications WHERE job_id=? ORDER BY attempt DESC LIMIT 1", (job_id,))

    if status in ("Applied", "Email Sent"):
        already = current and current["status"] in APPLIED_LIKE
        related = _related_applied(job)
        if (already or related) and not reapply:
            where = "this role" if already else f"the same role via {related['source']}"
            raise DuplicateApplicationError(
                f"You already applied to {where}. Confirm 're-apply' to record another attempt."
            )
        if already and reapply:
            current = None  # start a new attempt row
            attempt = (db.query_one("SELECT MAX(attempt) m FROM applications WHERE job_id=?", (job_id,)) or {"m": 0})["m"] + 1
        else:
            attempt = current["attempt"] if current else 1
    else:
        attempt = current["attempt"] if current else 1

    followup = None
    if status in ("Applied", "Email Sent"):
        followup = (datetime.now(timezone.utc) + timedelta(days=FOLLOWUP_DAYS)).replace(microsecond=0).isoformat()

    if current is None:
        history = [{"ts": now, "status": status, "note": notes or ""}]
        db.execute(
            "INSERT INTO applications(job_id,attempt,status,method,title,company,source,url,match_score,notes,confirmation,"
            "created_at,updated_at,applied_at,followup_due,history) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (job_id, attempt, status, method, job["title"], job["company"], job["source_name"], job["apply_url"] or job["url"],
             job["match_score"], notes, confirmation, now, now, now if status in APPLIED_LIKE else None, followup, json.dumps(history)),
        )
    else:
        history = db.jload(current["history"], []) + [{"ts": now, "status": status, "note": notes or ""}]
        db.execute(
            "UPDATE applications SET status=?, method=COALESCE(?,method), notes=COALESCE(?,notes), "
            "confirmation=COALESCE(?,confirmation), updated_at=?, applied_at=COALESCE(applied_at, ?), "
            "followup_due=COALESCE(?, followup_due), history=? WHERE id=?",
            (status, method, notes, confirmation, now, now if status in APPLIED_LIKE else None, followup,
             json.dumps(history), current["id"]),
        )
    if status in ("Shortlisted", "Ready to Apply"):
        db.execute("UPDATE jobs SET saved=1 WHERE id=?", (job_id,))
    log("tracker", "status", job_id=job_id, status=status, reapply=reapply)
    return get_for_job(job_id)  # type: ignore[return-value]


def list_applications(status: str | None = None) -> list[dict[str, Any]]:
    rows = db.query(
        "SELECT a.*, j.status AS job_status FROM applications a LEFT JOIN jobs j ON j.id=a.job_id "
        "WHERE a.attempt = (SELECT MAX(attempt) FROM applications b WHERE b.job_id=a.job_id) ORDER BY a.updated_at DESC"
    )
    out = [_decorate(r) for r in rows]
    if status:
        out = [r for r in out if r["status_display"] == status]
    return out


def delete_application(app_id: int) -> None:
    db.execute("DELETE FROM applications WHERE id=?", (app_id,))


EXPORT_COLS = ["updated_at", "status_display", "title", "company", "source", "url", "method", "match_score",
               "applied_at", "followup_due", "notes", "confirmation", "attempt", "job_status"]


def export_csv() -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([c.replace("_", " ").title() for c in EXPORT_COLS])
    for r in list_applications():
        w.writerow([r.get(c, "") for c in EXPORT_COLS])
    return buf.getvalue().encode("utf-8-sig")


def export_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    ws.append([c.replace("_", " ").title() for c in EXPORT_COLS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
    for r in list_applications():
        ws.append([r.get(c, "") for c in EXPORT_COLS])
    for i, width in enumerate([20, 14, 40, 28, 18, 50, 14, 10, 20, 20, 40, 20, 8, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    js = wb.create_sheet("Saved & top jobs")
    cols = ["match_score", "match_label", "title", "company", "location", "remote_type", "job_type", "source_name",
            "posted_at", "url", "status"]
    js.append([c.replace("_", " ").title() for c in cols])
    for cell in js[1]:
        cell.font = Font(bold=True)
    for r in db.query("SELECT * FROM jobs WHERE hidden=0 AND (saved=1 OR match_score>=65) ORDER BY match_score DESC LIMIT 500"):
        js.append([r.get(c) for c in cols])
    js.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
