"""Tracker, outreach rules / STOP, and API endpoints."""
import asyncio

import pytest
from fastapi.testclient import TestClient

from app import db, email_outreach as EO, job_tracker as JT
from app.settings_store import set_secret
from tests.test_pipeline import by_title, run_search

SENT = []


@pytest.fixture
def searched(api_mock, resume_profile):
    run_search()
    return run_search  # second run not needed by default


@pytest.fixture
def fake_sender(monkeypatch):
    SENT.clear()

    async def send(em):
        SENT.append(em)

    monkeypatch.setattr(EO, "SENDER_OVERRIDE", send)
    for k, v in {"SMTP_HOST": "smtp.example.com", "SMTP_USER": "me@example.com", "SMTP_PASSWORD": "app-pass"}.items():
        set_secret(k, v)
    # quiet hours off so tests are time-independent
    EO.save_settings({"quiet_hours": {"enabled": False}})
    return SENT


def rws_job():
    return by_title("Speech AI Evaluation")


# --------------------------------------------------------------- tracker ----
def test_tracker_status_flow_and_duplicate_protection(searched):
    jid = rws_job()["id"]
    JT.set_status(jid, "Shortlisted")
    JT.set_status(jid, "Applied", method="Official page", confirmation="REF-1")
    with pytest.raises(JT.DuplicateApplicationError):
        JT.set_status(jid, "Applied")
    again = JT.set_status(jid, "Applied", reapply=True)
    assert again["attempt"] == 2
    JT.set_status(jid, "Interview")
    apps = JT.list_applications()
    assert len(apps) == 1 and apps[0]["status"] == "Interview"
    assert [h["status"] for h in db.jload(db.query_one("SELECT history FROM applications WHERE attempt=1")["history"])] == [
        "Shortlisted", "Applied"]
    with pytest.raises(ValueError):
        JT.set_status(jid, "Bogus")


def test_tracker_exports(searched):
    JT.set_status(rws_job()["id"], "Applied")
    csv = JT.export_csv().decode("utf-8-sig")
    assert "Speech AI Evaluation" in csv
    from io import BytesIO
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(JT.export_xlsx()))
    assert wb.sheetnames == ["Applications", "Saved & top jobs"]
    assert wb["Applications"].max_row == 2


# -------------------------------------------------------------- outreach ----
def test_default_is_draft_only_and_never_sends(searched, fake_sender):
    msg = EO.create_draft(rws_job()["id"], "inquiry")
    assert msg["status"] == "draft"
    assert msg["recipient"] == "trainai.recruiting@rws-example.com"
    assert msg["recipient_source"] == "Published in job posting"
    asyncio.run(EO.process_queue())
    assert SENT == []
    with pytest.raises(EO.OutreachError):
        EO.approve(msg["id"])


def test_no_guessed_addresses(searched):
    with pytest.raises(EO.OutreachError, match="No email address"):
        EO.create_draft(by_title("Hindi Linguistic QA Tester")["id"])


def test_auto_mode_requires_rules(searched):
    with pytest.raises(EO.OutreachError, match="connect an email account"):
        EO.save_settings({"mode": "auto", "approved_categories": ["Hindi AI Trainer"]})


def test_review_mode_send_and_duplicate_outreach_blocked(searched, fake_sender):
    EO.save_settings({"mode": "review"})
    msg = EO.create_draft(rws_job()["id"], "inquiry")
    assert msg["status"] == "pending_review"
    EO.approve(msg["id"])
    res = asyncio.run(EO.process_queue())
    assert res["sent"] == 1 and len(SENT) == 1
    assert "Speech AI Evaluation" in SENT[0]["Subject"]
    assert JT.get_for_job(rws_job()["id"])["status"] == "Email Sent"
    # same job, same recipient again -> blocked
    m2 = EO.create_draft(rws_job()["id"], "inquiry")
    with pytest.raises(EO.OutreachError):
        EO.approve(m2["id"])


def test_auto_mode_respects_rules_and_caps(searched, fake_sender):
    EO.save_settings({"mode": "auto", "approved_categories": ["Audio and Speech Data Evaluator"], "min_score": 60,
                      "recipient_types": ["recruiter", "careers_inbox", "hiring_contact"], "max_per_day": 1,
                      "templates": ["application", "inquiry"]})
    session = db.query_one("SELECT id FROM search_sessions ORDER BY started_at DESC")["id"]
    db.execute("UPDATE jobs SET last_session_id=?", (session,))
    queued = EO.auto_outreach_for_session(session)
    assert queued == 1
    asyncio.run(EO.process_queue())
    assert len(SENT) == 1
    # Posting says "send your CV" -> application template chosen
    assert db.query_one("SELECT template FROM outreach WHERE status='sent'")["template"] == "application"


def test_emergency_stop_cancels_queue(searched, fake_sender):
    EO.save_settings({"mode": "review"})
    msg = EO.create_draft(rws_job()["id"], "inquiry")
    EO.approve(msg["id"])
    n = EO.emergency_stop()
    assert n == 1
    asyncio.run(EO.process_queue())
    assert SENT == []
    assert EO.get_message(msg["id"])["status"] == "cancelled"
    m2 = EO.create_draft(rws_job()["id"], "inquiry")
    with pytest.raises(EO.OutreachError, match="STOP"):
        EO.approve(m2["id"])
    EO.resume_sending()
    assert not EO.get_settings()["stopped"]


def test_opt_out_and_blocked_company(searched, fake_sender):
    EO.save_settings({"mode": "review"})
    EO.opt_out("trainai.recruiting@rws-example.com")
    msg = EO.create_draft(rws_job()["id"], "inquiry")
    with pytest.raises(EO.OutreachError, match="do-not-contact"):
        EO.approve(msg["id"])


def test_delivery_failures_trigger_auto_stop(searched, monkeypatch, fake_sender):
    async def boom(em):
        raise ConnectionError("smtp down")

    monkeypatch.setattr(EO, "SENDER_OVERRIDE", boom)
    EO.save_settings({"mode": "review", "max_per_company": 3, "stop_conditions": {"max_consecutive_failures": 1}})
    msg = EO.create_draft(rws_job()["id"], "inquiry")
    EO.approve(msg["id"])
    res = asyncio.run(EO.process_queue())
    assert res["failed"] == 1
    assert EO.get_message(msg["id"])["status"] == "failed"
    assert EO.get_settings()["stopped"] is True


def test_provisional_profile_drafts_cannot_send(api_mock, fake_sender):
    run_search()  # no resume uploaded -> provisional profile
    EO.save_settings({"mode": "review"})
    msg = EO.create_draft(rws_job()["id"], "inquiry")
    assert "PROVISIONAL" in msg["body"]
    with pytest.raises(EO.OutreachError, match="provisional"):
        EO.approve(msg["id"])


def test_quiet_hours():
    from datetime import datetime, timezone

    q = {"enabled": True, "start": "21:00", "end": "09:00", "tz": "Asia/Kolkata"}
    assert EO.in_quiet_hours(q, datetime(2026, 10, 27, 17, 0, tzinfo=timezone.utc))  # 22:30 IST
    assert not EO.in_quiet_hours(q, datetime(2026, 10, 27, 6, 0, tzinfo=timezone.utc))  # 11:30 IST


def test_audit_log_has_no_secrets(searched, fake_sender):
    rows = db.query("SELECT detail FROM audit_log")
    assert all("app-pass" not in (r["detail"] or "") for r in rows)


# ------------------------------------------------------------------- API ----
def test_api_end_to_end(api_mock, tmp_path):
    from app.main import app

    with TestClient(app) as c:
        assert c.get("/").status_code == 200
        st = c.get("/api/status").json()
        assert st["profile"]["verified"] is False and st["search"] is None  # idle on start
        resume = (tmp_path.parent / "x.txt")
        from tests.conftest import FIX

        r = c.post("/api/resume", files={"file": ("cv.txt", (FIX / "sample_resume.txt").read_bytes(), "text/plain")})
        assert r.status_code == 200 and r.json()["profile"]["name"] == "Priya Sharma"
        bad = c.post("/api/resume", files={"file": ("cv.exe", b"MZ" * 100, "application/octet-stream")})
        assert bad.status_code == 422
        run_search()  # direct pipeline run (TestClient loop + respx)
        jobs = c.get("/api/jobs").json()
        assert jobs["total"] >= 5 and jobs["items"][0]["score"] >= jobs["items"][-1]["score"]
        assert c.get("/api/jobs?quick=ai").json()["total"] >= 1
        inel = c.get("/api/jobs?include_ineligible=true").json()["total"]
        assert inel > jobs["total"]
        jid = jobs["items"][0]["id"]
        d = c.get(f"/api/jobs/{jid}").json()
        assert d["match"]["components"]["skills"]["max"] == 30
        assert c.get(f"/api/jobs/{jid}/cover-letter").json()["text"].startswith("Dear")
        assert c.get(f"/api/jobs/{jid}/tailored-resume.docx").status_code == 200
        fa = c.get(f"/api/jobs/{jid}/form-answers").json()
        assert any(n["field"] == "salary_expectation" for n in fa["needs_your_input"])
        assert c.post(f"/api/jobs/{jid}/status", json={"status": "Applied"}).status_code == 200
        assert c.post(f"/api/jobs/{jid}/status", json={"status": "Applied"}).status_code == 409
        assert c.get("/api/export/applications.csv").status_code == 200
        assert c.post("/api/command", json={"command": "NOPE"}).status_code == 400
        assert c.put("/api/sources/careers", json={"config": {"pages": ["http://127.0.0.1/admin"]}}).status_code == 422
        r = c.post("/api/sources", json={"kind": "greenhouse", "name": "Acme Loc", "config": {"boards": [{"token": "acmeloc"}]}})
        assert r.status_code == 200
        assert any(s["id"] == r.json()["id"] for s in c.get("/api/sources").json())
        assert c.post("/api/stop").json()["outreach_stopped"] is True
        assert c.put("/api/secrets", json={"name": "NOT_A_KEY", "value": "x"}).status_code == 422
        assert c.put("/api/settings", json={"weights": {"skills": -1}}).status_code == 422
        assert c.post("/api/jobs/rescore").json()["rescored"] >= 5
        assert c.delete("/api/resume").status_code == 200
        assert c.get("/api/status").json()["profile"]["verified"] is False


def test_go_via_api_runs_in_background_and_reports(api_mock, resume_profile, monkeypatch):
    import time

    from app import http_client
    from app.main import app

    # no throttling / caching in tests
    orig = http_client.HttpClient.__init__
    monkeypatch.setattr(http_client.HttpClient, "__init__",
                        lambda self, **kw: orig(self, **{**kw, "respect_intervals": False, "cache_ttl": 0, "backoff_base": 0.01}))
    with TestClient(app) as c:
        r = c.post("/api/command", json={"command": "GO"})
        assert r.status_code == 200 and r.json()["status"] == "running"
        assert c.post("/api/search", json={}).status_code == 409  # one search at a time
        for _ in range(100):
            st = c.get("/api/search/status").json()
            if st["status"] != "running":
                break
            time.sleep(0.1)
        assert st["status"] == "completed" and st["stage"] == "Completed"
        assert st["sources"]["lever"]["state"] == "done"
        assert st["stats"]["eligible_matches"] >= 5
        assert c.get("/api/search/sessions").json()[0]["status"] == "completed"
        assert c.get("/api/jobs?session_only=true").json()["total"] >= 5
